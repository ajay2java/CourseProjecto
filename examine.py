# import csv

from dict import general, accounting, business_analytics, comp_math_fin, econ, entre, enviro_sust, finance, global_reg_studies, hist_pol, iden_diver, info_tech, int_business, jcs, leadership, legal_studies, lit_visual, man_fin_pl, marketing, operations_mgmt, quant_m, real_estate, retail_scm, social_culture, strat_cult, tech_entr, tech_entr_des
# from app import f 

from openpyxl import load_workbook

# workbook = load_workbook(filename="data/Course_template.xlsx")


# workbook = load_workbook(filename=f)


# workbook.sheetnames

# sheet = workbook.active

num = ['0', '1', '2', '3', '4', '5', '6', '7', '8', '9']


class Concentration:
    
    def __init__(self, name, pre_req, course_list=None):
        ###
        self.name = name
        self.total_creds = int(pre_req)
        self.fMe = 0
        self.std = int(pre_req)
        self.adv_exp = 0
        self.adv_lib = 0
        self.adv_libe = 0
        self.gen_e = 0

        self.acc_req = 0
        self.acc_ele = 0

        self.ana_dm = 0
        self.ana_ad = 0
        self.ana_e = 0

        self.compr = 0
        self.compe1 = 0
        self.compe2 = 0

        self.econr = 0
        self.econe = 0

        self.entr = 0
        self.ente = 0

        self.env = 0

        self.fin = 0

        self.grs1 = 0
        self.grs2 = 0

        self.hist = 0
        self.pol = 0

        self.idr = 0
        self.ide = 0

        self.itm_c = 0
        self.itm_e = 0

        self.ibe_r = 0
        self.ibe_e = 0

        self.jcs_phil = 0
        self.jcs_e = 0

        self.lead_r = 0
        self.lead_r1 = 0
        self.lead_r2 = 0

        self.ls1 = 0
        self.ls2 = 0

        self.lva_r = 0

        self.mf_r = 0
        self.mf_e1 = 0
        self.mf_e2 = 0

        self.mkr = 0
        self.mke1 = 0
        self.mke2 = 0

        self.omr = 0
        self.ome = 0

        self.qm = 0

        self.re = 0

        self.rscr = 0
        self.rsce = 0

        self.sci = 0
        self.sca = 0

        self.sc_r = 0
        self.sc_e = 0

        self.ter = 0
        self.tee = 0

        self.ted_r1 = 0
        self.ted_r2 = 0
        self.ted_r3 = 0

        if course_list is None:
            self.courses = []
        else:
            self.courses = course_list
    
    def __str__(self):
        t = [f"{self.name} College Course Status:"]
        if self.total_creds >= 128:
            s = "Total Credits Status: ✅"
            t.append(s)
        else:
            s = "Total Credits Status: ❌"
            t.append(s)
        
        if self.fMe >= 2:
            s = "FME/EPS Requirement: ✅"
            t.append(s)
        else:
            s = "FME/EPS Requirements: ❌"
            t.append(s)

        if self.std >= 20:
            s = "Standard Requirements: ✅"
            t.append(s)
        else:
            s = "Standard Requirements: ❌"
            t.append(s)

        if self.adv_exp >= 1:
            s = "Advanced Experiential: ✅"
            t.append(s)
        else:
            s = "Advanced Experiential: ❌"
            t.append(s)

        if self.adv_lib >= 4: # num in credits
            s = "Advanced Liberal Arts: ✅"
            t.append(s)
        else:
            s = "Advanced Liberal Arts: ❌"
            t.append(s)
        
        if self.adv_libe >= 12:
            s = "Advanced Liberal Arts Elective: ✅"
            t.append(s)
        else:
            s = "Advanced Liberal Arts Elective: ❌"
            t.append(s)
        
        if self.gen_e >= 20:
            s = "General Elective: ✅"
            t.append(s)
        else:
            s = "General Elective: ❌"
            t.append(s)
        
        if self.acc_req == 3 and self.acc_ele >= 1:
            s = "Accounting Concentration: ✅"
            t.append(s)

        if self.ana_dm >= 1 and self.ana_ad >= 1 and self.ana_e >= 2:
            s = "Business Analytics Concentration: ✅"
            t.append(s)
        # repeat for other concentrations
        
        if self.compr == 1 and self.compe1 >= 2 and self.compe2 >= 1:
            s = "Computational and Mathematical Finance Concentration: ✅"
            t.append(s)
        # return "\n".join(t)

        if self.econr >= 1 and self.econe >= 3:
            s = "Economics Concentration: ✅"
            t.append(s)

        if self.entr == 1 and self.ente >= 3:
            s = "Entrepreneurship Concentration: ✅"
            t.append(s)

        if self.env >= 4:
            s = "Environmental Sustainability Concentration: ✅"
            t.append(s)
        
        if self.fin >= 4:
            s = "Finance Concentration: ✅"
            t.append(s)

        if self.grs1 >= 1 and self.grs2 >= 1 and self.grs1 + self.grs2 >= 4:
            s = "Global and Regional Studies Concentration: ✅"
            t.append(s)

        if self.hist >= 1 and self.pol >= 1 and self.hist + self.pol >= 4:
            s = "Historical and Political Studies Concentration: ✅"
            t.append(s)

        if self.idr >= 2 and self.idr + self.ide >= 4:
            s = "Identity and Diversity Concentration: ✅"
            t.append(s)
        
        if self.itm_c >= 4 and self.itm_c + self.itm_e >= 16:
            s = "Information Technology Management Concentration: ✅"
            t.append(s)
        
        #check on this one
        if self.ibe_r >= 3 and self.ibe_e >= 1:
            s = "International Business Environment Concentration: ✅"
            t.append(s)
        
        #check on this one
        if self.jcs_phil + self.jcs_e >= 4:
            s = "Justice, Citizenship, and Social Responsibility Concentration: ✅"
            t.append(s)

        if self.lead_r == 1 and self.lead_r1 >= 2 and self.lead_r2 >= 1:
            s = "Leadership Concentration: ✅"
            t.append(s)

        if self.ls1 >= 3 and self.ls1 + self.ls2 >= 4:
            s = "Legal Studies Concentration: ✅"
            t.append(s)

        #check this one
        if self.lva_r >= 16:
            s = "Literary and Visual Arts Concentration: ✅"
            t.append(s)

        if self.mf_r == 2 and self.mf_e1 >= 1 and self.mf_e2 >= 1:
            s = "Managerial Financial Planning and Analysis Concentration: ✅"
            t.append(s)

        if self.mkr == 1 and self.mke1 >= 1 and self.mke2 >= 2:
            s = "Marketing Concentration: ✅"
            t.append(s)

        if self.omr == 1 and self.ome >= 3:
            s = "Operations Management Concentration: ✅"
            t.append(s)
        
        if self.qm >= 4:
            s = "Quantitative Methods Concentration: ✅"
            t.append(s)

        if self.re >= 4:
            s = "Real Estate Concentration: ✅"
            t.append(s)

        if self.rscr == 2 and self.rsce >= 2:
            s = "Retail Supply Chain Management Concentration: ✅"
            t.append(s)

        if self.sci >= 2 and self.sca >= 2:
            s = "Social and Cultural Studies Concentration: ✅"
            t.append(s)

        if self.sc_r >= 2 and self.sc_e >= 2:
            s = "Strategy & Consulting Concentration: ✅"
            t.append(s)

        if self.ter >= 2 and self.tee >= 2:
            s = "Tech Entrepreneurship Concentration: ✅"
            t.append(s)

        if self.ted_r1 == 1 and self.ted_r2 >= 1 and self.ted_r3 >= 2:
            s = "Technology, Entrepreneurship, and Design Concentration: ✅"
            t.append(s)
        
        return t
            
    
    def num_credits(self):
        for a in self.courses:
            if a in general['One Credit']:
                self.total_creds += 1
            elif a in general['Two Credit']:
                self.total_creds += 2
            elif a in general['Three Credit']:
                self.total_creds += 3
            elif a in general['Six Credit']:
                self.total_creds += 6
            elif a in general['Seventeen Credit']:
                self.total_creds += 17
            else:
                self.total_creds += 4
    
    
    def get_num_creds(self, course):
        if course in general['One Credit']:
            return 1
        elif course in general['Two Credit']:
            return 2
        elif course in general['Three Credit']:
            return 3
        elif course in general['Six Credit']:
            return 6
        elif course in general['Seventeen Credit']:
            return 17
        else:
            return 4

    def FME(self):
        for a in self.courses:
            if a in general['FME']:
                self.fMe += 1

    def standard(self):
        for a in self.courses:
            if a in general['Standard']:
                self.std += 1

    def adv_exper(self):
        for a in self.courses:
            if a in general['Adv_Experiential']:
                self.adv_exp += 1

    def adv_libarts(self):
        store = ""
        for a in self.courses:
            for i in range(len(a)):
                if a[i] in num:
                    store += a[i]
            if len(store) == 4 and '46' == store[0:2]:
                self.adv_lib += self.get_num_creds(a)
            store = ""

    def adv_libarts_elec(self):
        store = ""
        for a in self.courses:
            # print(a)
            for i in range(len(a)):
                if a[i] in num:
                    store += a[i]
            if len(store) == 4 and store[1] == '6':
                self.adv_libe += self.get_num_creds(a)
            store = ""

    def gen_elective(self):
        store = ""
        for a in self.courses:
            # print(a)
            for i in range(len(a)):
                if a[i] in num:
                    store += a[i]
            if len(store) == 4 and (store[1] == '1' or store[1] == '2' or store[1] == '5' or store[1] == '6'):
                self.gen_e += self.get_num_creds(a)
            store = ""

    def accounting(self):
        for a in self.courses:
            if a in accounting['Required']:
                self.acc_req += 1
            elif a in accounting['Elective']:
                self.acc_ele += 1



    def analytics(self):
        for a in self.courses:
            if a in business_analytics['Data Management and Programming Concepts']:
                if self.ana_dm == 1:
                    self.ana_e += self.get_num_creds(a)
                else:
                    self.ana_dm += 1
            elif a in business_analytics['Advanced Data and Decision Modeling']:
                if self.ana_ad == 1:
                    self.ana_e += self.get_num_creds(a)
                else:
                    self.ana_ad += 1
            elif a in business_analytics['Electives']:
                self.ana_e += self.get_num_creds(a)
            else:
                continue

    def computational(self):
        for a in self.courses:
            if a in comp_math_fin['Required']:
                self.compr += 1
            elif a in comp_math_fin['Elective_1']:
                self.compe1 += 1
            elif a in comp_math_fin['Elective_2']:
                self.compe2 += 1

    def economics(self):
        for a in self.courses:
            if a in econ['Required']:
                self.econr += 1
            elif a in econ['Elective']:
                self.econe += 1

    def entrepren(self):
        for a in self.courses:
            if a in entre['Required']:
                self.entr += 1
            elif a in entre['Elective']:
                self.ente += 1
    
    def enviro_sustainability(self):
        for a in self.courses:
            if a in enviro_sust['General']:
                self.env += 1

    def financee(self):
        for a in self.courses:
            if a in finance['General']:
                self.fin += 1
        
    def global_regional_studies(self):
        for a in self.courses:
            if a in global_reg_studies['Global']:
                self.grs1 += 1
            elif a in global_reg_studies['Regional']:
                self.grs2 += 1

    def historical_political(self):
        for a in self.courses:
            if a in hist_pol['Historical']:
                if self.hist == 1:
                    if a[3] in num:
                        if int(a[3]) >= 3:
                           self.hist += 1 
                else:
                    self.hist += 1
            elif a in hist_pol['Political']:
                if self.pol == 1:
                    if a[3] in num:
                        if int(a[3]) >= 3:
                           self.pol += 1 
                else:
                    self.pol += 1
    
    def identity_diversity(self):
        for a in self.courses:
            if a in iden_diver['Required']:
                self.idr += 1
            elif a in iden_diver['Elective']:
                self.ide += 1
    
    def info_tech_management(self):
        for a in self.courses:
            if a in info_tech['Creator']:
                self.itm_c += self.get_num_creds(a)
            elif a in info_tech['Elective']:
                self.itm_e += self.get_num_creds(a)

    def inter_business_enviro(self):
        for a in self.courses:
            if a in int_business['Required']:
                self.ibe_r += 1
            elif a in int_business['Elective']:
                self.ibe_e += 1

    def just_citi_sr(self):
        for a in self.courses:
            if a in jcs['Philosophy']:
                self.jcs_phil += 1
            elif a in jcs['Elective']:
                self.jcs_e += 1

    def lead(self):
        for a in self.courses:
            if a in leadership['Required']:
                self.lead_r += 1
            elif a in leadership['Req2']:
                self.lead_r1 += 1
            elif a in leadership['Req3']:
                self.lead_r2 += 1

    def legal_study(self):
        for a in self.courses:
            if a in legal_studies['Part1']:
                self.ls1 += 1
            elif a in legal_studies['Part2']:
                self.ls2 += 1

    def lva(self):
        for a in self.courses:
            if a in lit_visual['Required']:
                self.lva_r += self.get_num_creds(a)

    def mfpa(self):
        for a in self.courses:
            if a in man_fin_pl['Required']:
                self.mf_r += 1
            elif a in man_fin_pl['Elective1']:
                self.mf_e1 += 1
            elif a in man_fin_pl['Elective2']:
                self.mf_e2 += 1

    def market(self):
        for a in self.courses:
            if a in marketing['Required']:
                self.mkr += 1
            elif a in marketing['Elective1']:
                if self.mke1 == 1:
                    self.mke2 += 1
                else:
                    self.mke1 += 1
            elif a in marketing['Elective2']:
                self.mke2 += 1
    
    def operations_management(self):
        for a in self.courses:
            if a in operations_mgmt['Required']:
                self.omr += 1
            elif a in operations_mgmt['Elective']:
                self.ome += 1

    def quant_methods(self):
        for a in self.courses:
            if a in quant_m['Required']:
                self.qm += 1
    
    def real_e(self):
        for a in self.courses:
            if a in real_estate['Required']:
                self.re += 1

    def retail_supplychain(self):
        for a in self.courses:
            if a in retail_scm['Required']:
                self.rscr += 1
            elif a in retail_scm['Elective']:
                self.rsce += 1

    def social_cultural_studies(self):
        for a in self.courses:
            if a in social_culture['Intermediate']:
                self.sci += 1
            elif a in social_culture['Advanced']:
                self.sca += 1
    
    def strategy_consulting(self):
        for a in self.courses:
            if a in strat_cult['Required']:
                self.sc_r += 1
            elif a in strat_cult['Elective']:
                self.sc_e += 1
    
    def tech_entrepreneurship(self):
        for a in self.courses:
            if a in tech_entr['Required']:
                self.ter += 1
            elif a in tech_entr['Elective']:
                self.tee += 1
    
    def tech_entre_design(self):
        for a in self.courses:
            if a in tech_entr_des['Required']:
                self.ted_r1 += 1
            elif a in tech_entr_des['Req2']:
                self.ted_r2 += 1
            elif a in tech_entr_des['Elective']:
                self.ted_r3 += 1


def store1(workbook):

    # workbook = load_workbook(filename="data/{f}")
    # workbook = load_workbook(filename=f)

    workbook.sheetnames

    sheet = workbook.active

    ### code that takes all courses from excel and puts them into a list##
    course_list = [sheet["A2"].value, sheet["A3"].value, sheet["A4"].value, sheet["A5"].value, sheet["A6"].value,
                sheet["A11"].value, sheet["A12"].value, sheet["A13"].value, sheet["A14"].value, sheet["A15"].value,

                sheet["E2"].value, sheet["E3"].value, sheet["E4"].value, sheet["E5"].value, sheet["E6"].value,
                sheet["E11"].value, sheet["E12"].value, sheet["E13"].value, sheet["E14"].value, sheet["E15"].value,

                sheet["I2"].value, sheet["I3"].value, sheet["I4"].value, sheet["I5"].value, sheet["I6"].value,
                sheet["I11"].value, sheet["I12"].value, sheet["I13"].value, sheet["I14"].value, sheet["I15"].value,
                
                sheet["M2"].value, sheet["M3"].value, sheet["M4"].value, sheet["M5"].value, sheet["M6"].value,
                sheet["M11"].value, sheet["M12"].value, sheet["M13"].value, sheet["M14"].value, sheet["M15"].value]




    for z in course_list:
        if z == sheet["B3"].value:
            course_list.remove(z)


    if sheet["C18"].value == sheet["B3"].value:
        id = "Student"  
    else:
        id = sheet["C18"].value

    #Put an int check on this!
    # pre_req = input("How many credits from pre-Babson do you have, which counts? ")

    if sheet["H18"].value == sheet["B3"].value:
        pre_req = 0
    else:
        pre_req = int(sheet["H18"].value)

    Object1 = Concentration(id, pre_req, course_list)
    Object1.num_credits()
    Object1.FME()
    Object1.standard()
    Object1.adv_exper()
    Object1.adv_libarts()
    Object1.adv_libarts_elec()
    Object1.gen_elective()
    
    Object1.accounting()
    Object1.analytics()
    Object1.computational()
    Object1.economics()
    Object1.entrepren()
    Object1.enviro_sustainability()
    Object1.financee()
    Object1.global_regional_studies()
    Object1.historical_political()
    Object1.identity_diversity()
    Object1.info_tech_management()
    Object1.inter_business_enviro()
    Object1.just_citi_sr()
    Object1.lead()
    Object1.legal_study()
    Object1.lva()
    Object1.mfpa()
    Object1.market()
    Object1.operations_management()
    Object1.quant_methods()
    Object1.real_e()
    Object1.retail_supplychain()
    Object1.social_cultural_studies()
    Object1.strategy_consulting()
    Object1.tech_entrepreneurship()
    Object1.tech_entre_design()


    # print(Ajay)
    # return(str(Object1))
    return(Object1.__str__())


# store1("data/Course_template.xlsx")
